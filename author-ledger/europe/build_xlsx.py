"""Build the Europe author-lead workbook (READY / NEEDS REVIEW / DUPLICATES / COUNTRY PROGRESS)."""
import collections, datetime, glob, json, os, re, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
HEAD = ["Author Full Name", "Country", "Book Title", "Genre", "Publication Date", "Author Email",
        "Email Type (Direct / Agent / Publisher)", "Email Verification Source URL", "Author Website",
        "Publisher", "Book Source URL", "Recent Author Activity", "Personalization Detail",
        "Commercial Fit", "Outreach Language", "Duplicate Status", "Qualification Status",
        "Research Date", "Author Spotlight Workflow 2 Notes"]
KEYS = ["author_full_name", "country", "book_title", "genre", "publication_date", "author_email",
        "email_type", "email_source_url", "author_website", "publisher", "book_source_url",
        "recent_activity", "personalization_detail", "commercial_fit", "outreach_language",
        "duplicate_status", "qualification_status", "research_date", "workflow2_notes"]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
GENERIC = re.compile(r"^(info|orders|order|sales|submissions?|hello|contact|office|admin|webmaster)@", re.I)
TODAY = datetime.date.today().isoformat()

def norm(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", s)

master = {norm(l) for l in open(os.path.join(HERE, "work", "EXISTING_MASTER_NAMES.txt")) if l.strip()}
ready, review, dups, seen, seen_direct = [], [], [], {}, set()
FIX = {}
for fp in sorted(glob.glob(os.path.join(HERE, "work", "zz_fixes*.jsonl"))):
    for line in open(fp, encoding="utf-8"):
        if line.strip():
            try:
                f = json.loads(line); FIX[norm(f.get("author_full_name", ""))] = f
            except Exception:
                pass
for path in sorted(p for p in glob.glob(os.path.join(HERE, "work", "*.jsonl")) if not os.path.basename(p).startswith("zz_fixes")):
    for line in open(path, encoding="utf-8"):
        if not line.strip():
            continue
        try:
            r = json.loads(line)
        except Exception:
            continue
        if norm(r.get("author_full_name", "")) in FIX:
            r = FIX.pop(norm(r.get("author_full_name", "")))
        if str(r.get("qualification_status", "")).upper() == "EXCLUDED":
            continue
        r = {k: (r.get(k) or "") for k in KEYS + ["review_reason"]}
        r["research_date"] = r["research_date"] or TODAY
        k = norm(r["author_full_name"])
        if not k:
            continue
        if k in master:
            r["duplicate_status"] = "DUPLICATE — already in Author Ledger master"
            r["qualification_status"] = "DUPLICATE"; dups.append(r); continue
        if k in seen:
            r["duplicate_status"] = f"DUPLICATE — also found in {seen[k]}"
            r["qualification_status"] = "DUPLICATE"; dups.append(r); continue
        seen[k] = os.path.basename(path)
        em = r["author_email"].strip()
        et = r["email_type"].strip().capitalize()
        problems = []
        if not em or not EMAIL_RE.match(em): problems.append("no valid email")
        elif GENERIC.match(em) and r["email_type"].strip().capitalize() != "Direct": problems.append("generic company inbox (not acceptable)")
        if em and not r["email_source_url"]: problems.append("no email source URL")
        if et not in ("Direct", "Agent", "Publisher"): problems.append("email type missing")
        if et == "Direct" and em.lower() in seen_direct: problems.append("direct email shared with another author")
        if not re.search(r"202[5-7]", r["publication_date"]): problems.append("no 2025–2027 publication date")
        if not r["book_title"]: problems.append("no book title")
        r["email_type"] = et
        r["duplicate_status"] = "New"
        if r["qualification_status"].upper() == "READY" and not problems:
            r["qualification_status"] = "READY"; ready.append(r)
            if et == "Direct": seen_direct.add(em.lower())
        else:
            reason = "; ".join(problems + ([r["review_reason"]] if r["review_reason"] else []))
            r["qualification_status"] = "NEEDS REVIEW"
            r["workflow2_notes"] = (r["workflow2_notes"] + " | Review: " + reason).strip(" |")
            review.append(r)

wb = Workbook()
def sheet(ws, rows):
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in sorted(rows, key=lambda x: (x["country"], x["author_full_name"])):
        ws.append([r[k] for k in KEYS])
    ws.freeze_panes = "B2"
    widths = [24, 14, 30, 16, 12, 32, 14, 40, 30, 22, 40, 40, 50, 22, 14, 22, 14, 12, 50]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = "READY"; sheet(ws, ready)
sheet(wb.create_sheet("NEEDS REVIEW"), review)
sheet(wb.create_sheet("DUPLICATES"), dups)
cp = wb.create_sheet("COUNTRY PROGRESS")
cp.append(["Country", "READY (qualified)", "Direct emails", "Agent emails", "Publisher emails", "Needs review", "Duplicates rejected"])
for c in cp[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
countries = sorted({r["country"] for r in ready + review + dups if r["country"]})
for c in countries:
    rr = [r for r in ready if r["country"] == c]
    cp.append([c, len(rr), sum(r["email_type"] == "Direct" for r in rr), sum(r["email_type"] == "Agent" for r in rr),
               sum(r["email_type"] == "Publisher" for r in rr), sum(r["country"] == c for r in review),
               sum(r["country"] == c for r in dups)])
cp.append(["TOTAL", len(ready), sum(r["email_type"] == "Direct" for r in ready), sum(r["email_type"] == "Agent" for r in ready),
           sum(r["email_type"] == "Publisher" for r in ready), len(review), len(dups)])
for c in cp[cp.max_row]: c.font = Font(bold=True)
cp.column_dimensions["A"].width = 24
out = os.path.join(HERE, "Author_Ledger_Europe_Leads.xlsx"); wb.save(out)
print(f"READY {len(ready)} | NEEDS REVIEW {len(review)} | DUPLICATES {len(dups)} -> {out}")
print(collections.Counter(r["country"] for r in ready).most_common())
print(collections.Counter(r["email_type"] for r in ready))
