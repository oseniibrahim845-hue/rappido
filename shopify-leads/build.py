"""Build Shopify lead deliverables from work/*.jsonl (research rows) and work/zz_fix*.jsonl (overrides).

Outputs:
  Shopify_Inventory_Order_Automation_700_Leads.xlsx  (sheet Shopify_Leads = qualified; Needs_Review; Duplicates_Removed)
  Shopify_Research_Summary.xlsx
"""
import collections, datetime, glob, json, os, re
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
TODAY = datetime.date.today().isoformat()
COLS = ["lead_id", "store_name", "website", "country", "product_niche", "owner_name", "first_name", "owner_role",
        "email", "email_type", "shopify_verified", "shopify_verification_source", "product_or_collection",
        "personalized_store_observation", "inventory_order_automation_opportunity", "opportunity_reasoning",
        "personalized_subject", "personalized_email_body", "research_source_url", "contact_source_url",
        "email_verification_status", "personalization_status", "automation_action", "send_status", "date_sent",
        "gmail_thread_id", "reply_status", "do_not_contact", "last_error", "research_date", "confidence_score"]
TARGET = {"United States": 250, "United Kingdom": 150, "Canada": 100, "Australia": 100, "Germany": 100}
PREFIX = {"United States": "US", "United Kingdom": "UK", "Canada": "CA", "Australia": "AU", "Germany": "DE"}
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
BAD_DOMAINS = ("amazon.", "etsy.com", "ebay.", "gmail.com", "yahoo.", "hotmail.", "outlook.com")


def domain(u):
    u = (u or "").strip().lower()
    if not u.startswith("http"):
        u = "https://" + u
    d = urlparse(u).netloc
    return d[4:] if d.startswith("www.") else d


SENT = {}
_sl = os.path.join(os.path.dirname(os.path.abspath(__file__)), "work", "send_log.jsonl")
if os.path.exists(_sl):
    for _l in open(_sl, encoding="utf-8"):
        if _l.strip():
            _d = json.loads(_l); SENT[_d["email"].lower()] = _d
BOUNCED = {}
_bl = os.path.join(os.path.dirname(_sl), "bounces.jsonl")
if os.path.exists(_bl):
    for _l in open(_bl, encoding="utf-8"):
        if _l.strip():
            _d = json.loads(_l); BOUNCED[_d["email"].lower()] = _d


def words(body):
    body = re.split(r"\n\s*(Best regards|Viele Grüße|Mit freundlichen Grüßen),", body)[0]
    return len(re.findall(r"\b[\wÄÖÜäöüß'-]+\b", body))


NICHES = [("Pet products and accessories", r"\bpet|dog|cat\b|puppy|canine|feline"),
          ("Baby and children's products", r"baby|kid|child|toddler|nursery|infant|toy"),
          ("Jewelry and accessories", r"jewel|ring|necklace|earring|watch|bracelet|schmuck"),
          ("Skincare and beauty products", r"skin|beauty|cosmetic|makeup|hair|fragrance|perfume|soap|self-tan|lip|nail|kosmetik"),
          ("Automotive accessories", r"auto|car care|vehicle|truck|off-?road|motorcycle|moto|jeep|detail"),
          ("Electronics and accessories", r"electronic|phone|case|charger|cable|audio|laptop|gadget|tech|magsafe"),
          ("Fitness and sports equipment", r"fitness|gym|kettlebell|sport|cycl|bike|outdoor|camp|hammock|hik|climb|yoga|weight"),
          ("Specialty food and beverage brands", r"coffee|tea\b|teas\b|chocolate|sauce|snack|food|spice|honey|beverage|wine|candy|roast"),
          ("Health and wellness products", r"wellness|supplement|herbal|vitamin|health|sleep|cbd|gummies"),
          ("Kitchen and household products", r"kitchen|cook|skillet|cast iron|clean|laundry|household|storage|cutlery"),
          ("Home decor and furniture", r"home|decor|furniture|candle|rug|bedding|lighting|interior|wohn|pillow|throw|textile"),
          ("Fashion and clothing brands", r"fashion|cloth|apparel|wear|swim|denim|dress|shirt|boutique|shoe|footwear|hat|mode|bag")]


def niche(n, p=""):
    t = f"{n} {p}".lower()
    for name, pat in NICHES:
        if re.search(pat, str(n).lower()):
            return name
    for name, pat in NICHES:
        if re.search(pat, t):
            return name
    return n or "Other"


FIX = {}
for fp in sorted(glob.glob(os.path.join(HERE, "work", "zz_fix*.jsonl"))):
    for line in open(fp, encoding="utf-8"):
        if line.strip():
            f = json.loads(line); FIX[domain(f.get("website"))] = f

qual, review, dups = [], [], []
seen_dom, seen_mail = {}, {}
for path in sorted(glob.glob(os.path.join(HERE, "work", "*.jsonl"))):
    if os.path.basename(path).startswith("zz_"):
        continue
    for n, line in enumerate(open(path, encoding="utf-8")):
        if not line.strip():
            continue
        try:
            r = json.loads(line)
        except Exception:
            continue
        d = domain(r.get("website"))
        if d in FIX:
            r = FIX.pop(d)
        if str(r.get("personalization_status", "")).upper() == "EXCLUDED" or not d:
            continue
        r = {k: ("" if r.get(k) is None else r.get(k)) for k in COLS + ["review_reason"]}
        r["website"] = "https://" + d
        r["product_niche"] = niche(r["product_niche"], r["product_or_collection"])
        em = str(r["email"]).strip()
        src = os.path.basename(path)
        if d in seen_dom or (em and em.lower() in seen_mail):
            r["last_error"] = f"Duplicate of {seen_dom.get(d) or seen_mail.get(em.lower())}"
            dups.append(r); continue
        seen_dom[d] = src
        if em:
            seen_mail[em.lower()] = src
        probs = []
        if r["country"] not in TARGET: probs.append("country outside target")
        if str(r["shopify_verified"]).upper() != "YES" or not r["shopify_verification_source"]: probs.append("Shopify not verified")
        if not em or not EMAIL_RE.match(em): probs.append("no valid email")
        elif any(b in em.lower().split("@")[1] for b in BAD_DOMAINS) and r["email_type"] != "Decision-maker":
            probs.append("email domain not the business")
        if not r["contact_source_url"]: probs.append("no contact source URL")
        if not r["research_source_url"]: probs.append("no research source URL")
        for k in ("product_or_collection", "personalized_store_observation", "inventory_order_automation_opportunity",
                  "personalized_subject", "personalized_email_body"):
            if not str(r[k]).strip(): probs.append(f"missing {k}")
        body = str(r["personalized_email_body"])
        if body:
            w = words(body)
            if not 75 <= w <= 120: probs.append(f"email body {w} words (need 75–120)")
            if re.search(r"\[[^\]]*\]|\{[^}]*\}|<[^>]*>|XXX|TBD", body + r["personalized_subject"]): probs.append("placeholder left in email")
            if "Oseni Ibrahim" not in body: probs.append("sign-off missing")
            if not body.rstrip().split("Oseni")[0].strip().endswith(("?", ",")) and "?" not in body: probs.append("no closing question")
        r["email_type"] = r["email_type"] or ("General inbox" if em.split("@")[0].lower() in ("hello", "info", "support", "contact", "orders", "team", "shop", "help", "care", "customerservice", "service", "kontakt") else "")
        r["email_verification_status"] = r["email_verification_status"] or ("PUBLICLY_LISTED" if r["contact_source_url"] else "UNVERIFIED")
        r["send_status"] = "NOT_SENT"; r["reply_status"] = "NO_REPLY"; r["do_not_contact"] = "FALSE"
        r["date_sent"] = r["gmail_thread_id"] = ""
        sent = SENT.get(em.lower())
        if sent:
            r["send_status"] = "SENT"; r["date_sent"] = sent["date_sent"]; r["gmail_thread_id"] = sent["thread_id"]
        r["research_date"] = r["research_date"] or TODAY
        try:
            r["confidence_score"] = int(r["confidence_score"])
        except Exception:
            r["confidence_score"] = ""
        if str(r["personalization_status"]).upper() == "READY" and not probs:
            r["personalization_status"] = "READY"; r["automation_action"] = "READY"; r["last_error"] = ""
            qual.append(r)
        else:
            r["personalization_status"] = "NEEDS_REVIEW"; r["automation_action"] = "HOLD"
            r["last_error"] = "; ".join(probs + ([r["review_reason"]] if r["review_reason"] else []))
            review.append(r)
        if em.lower() in BOUNCED:
            r["send_status"] = "BOUNCED"; r["do_not_contact"] = "TRUE"
            r["last_error"] = "; ".join(x for x in (r["last_error"], "Bounced: " + BOUNCED[em.lower()]["reason"]) if x)

# lead ids per country
cnt = collections.Counter()
for r in sorted(qual, key=lambda x: (list(TARGET).index(x["country"]), x["store_name"].lower())):
    cnt[r["country"]] += 1
    r["lead_id"] = f"{PREFIX[r['country']]}-{cnt[r['country']]:03d}"
for i, r in enumerate(review, 1):
    r["lead_id"] = f"REV-{i:03d}"
for i, r in enumerate(dups, 1):
    r["lead_id"] = f"DUP-{i:03d}"

HFILL = PatternFill("solid", fgColor="1F3A5F")
def sheet(ws, rows):
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HFILL; c.alignment = Alignment(vertical="top")
    for r in sorted(rows, key=lambda x: x["lead_id"]):
        ws.append([r[k] for k in COLS])
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVWXYZ", [9, 24, 30, 14, 18, 18, 12, 16, 30, 14, 9, 40, 30, 50, 45, 45, 34, 70, 40, 40, 16, 14, 12, 10, 10, 12, 10, 10, 30]):
        ws.column_dimensions[col].width = w
    for c in ("AA", "AB", "AC", "AD", "AE"):
        ws.column_dimensions[c].width = 12

wb = Workbook(); ws = wb.active; ws.title = "Shopify_Leads"; sheet(ws, qual)
sheet(wb.create_sheet("Needs_Review"), review)
sheet(wb.create_sheet("Duplicates_Removed"), dups)
OUT = os.path.join(HERE, "Shopify_Inventory_Order_Automation_700_Leads.xlsx"); wb.save(OUT)

# summary workbook
sw = Workbook(); s = sw.active; s.title = "Overview"
def hdr(ws, row):
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HFILL
hdr(s, ["Metric", "Value"])
gen = sum(r["email_type"] == "General inbox" for r in qual); dm = sum(r["email_type"] == "Decision-maker" for r in qual)
for k, v in [("Qualified leads (READY)", len(qual)), ("Target", 700), ("Verified business emails (publicly listed)", sum(r["email_verification_status"] == "PUBLICLY_LISTED" for r in qual)),
             ("Decision-maker emails", dm), ("General inbox emails", gen), ("Leads needing review", len(review)),
             ("Duplicate records removed", len(dups)), ("Named decision maker found", sum(bool(r["first_name"]) for r in qual)),
             ("Research date", TODAY)]:
    s.append([k, v])
s.column_dimensions["A"].width = 44; s.column_dimensions["B"].width = 14
c = sw.create_sheet("By_Country"); hdr(c, ["Country", "Qualified", "Target", "Remaining", "Decision-maker emails", "General inbox", "Needs review"])
for k, t in TARGET.items():
    q = [r for r in qual if r["country"] == k]
    c.append([k, len(q), t, max(0, t - len(q)), sum(r["email_type"] == "Decision-maker" for r in q), sum(r["email_type"] == "General inbox" for r in q), sum(r["country"] == k for r in review)])
c.append(["TOTAL", len(qual), 700, max(0, 700 - len(qual)), dm, gen, len(review)])
c.column_dimensions["A"].width = 18
nn = sw.create_sheet("By_Niche"); hdr(nn, ["Product niche", "Qualified"])
for k, v in collections.Counter(r["product_niche"] for r in qual).most_common():
    nn.append([k, v])
nn.column_dimensions["A"].width = 34
op = sw.create_sheet("Automation_Opportunities"); hdr(op, ["Opportunity theme", "Leads"])
THEMES = [("Variant (size/colour) low-stock alerts", r"variant|size|colou?r|shade"), ("Bundle/kit inventory sync", r"bundle|kit|set|box"),
          ("Subscription / repeat-order workflows", r"subscri|repeat|reorder|refill|replenish"), ("Supplier restock notifications", r"supplier|restock|purchase order|vendor"),
          ("Pre-order / launch / drop stock management", r"pre-?order|launch|drop|limited|waitlist|back-in-stock"), ("Multi-location / multi-channel sync", r"location|channel|wholesale|retail|marketplace|warehouse"),
          ("Order status / fulfilment notifications", r"status|fulfil|shipping|dispatch|tracking"), ("Inventory reporting", r"report|dashboard|forecast")]
tc = collections.Counter()
for r in qual:
    t = str(r["inventory_order_automation_opportunity"]).lower()
    hit = next((n for n, p in THEMES if re.search(p, t)), "Other")
    tc[hit] += 1
for k, v in tc.most_common():
    op.append([k, v])
op.column_dimensions["A"].width = 46
md = sw.create_sheet("Missing_Data"); hdr(md, ["Issue (Needs_Review rows)", "Count"])
mc = collections.Counter()
for r in review:
    for p in str(r["last_error"]).split("; "):
        mc[re.sub(r"\d+ words", "N words", p)[:80]] += 1
for k, v in mc.most_common(40):
    md.append([k, v])
md.column_dimensions["A"].width = 80
nm = sw.create_sheet("Missing_Fields_Qualified"); hdr(nm, ["Field empty among qualified leads", "Count"])
for k in ("owner_name", "first_name", "owner_role"):
    nm.append([k, sum(not r[k] for r in qual)])
nm.column_dimensions["A"].width = 40
SUM = os.path.join(HERE, "Shopify_Research_Summary.xlsx"); sw.save(SUM)
print(f"QUALIFIED {len(qual)} | NEEDS_REVIEW {len(review)} | DUPLICATES {len(dups)}")
print(dict(collections.Counter(r["country"] for r in qual)))
print("email types", dict(collections.Counter(r["email_type"] for r in qual)))
