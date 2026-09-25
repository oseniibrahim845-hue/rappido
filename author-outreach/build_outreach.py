import csv, re, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

OUT = "/home/user/rappido/author-outreach"
os.makedirs(OUT, exist_ok=True)

COLS = ["author_name","email","first_name","book_title","book_topic","personalization_detail",
"personalization_reason","personalized_opening","subject_line","email_body","research_sources",
"research_confidence","personalization_status","send_status","date_sent","country","has_audiobook",
"audiobook_check_sources","has_author_website","website_url","website_quality","has_book_trailer",
"book_trailer_url","has_amazon_a_plus","best_offer","offer_reason"]

SPOT = ("I would love to invite you to be featured in an Author Spotlight on The Author Ledger. "
        "The feature is designed to introduce readers to your work, the story behind your book, and the "
        "experiences or ideas that have shaped your journey as an author.")
SERV = ("Alongside our editorial features, we also support authors with professionally produced audiobook "
        "editions, cinematic book trailers, short promotional reels, author websites, press kits, Amazon A+ "
        "graphics, and book relaunch materials.")
SIGN = "Warm regards,\nPopoola Michael\nThe Author Ledger"
A_PLUS_NOTE = "Unclear"
AMZ = "Amazon (amazon.com search attempted, blocked by network policy, not verified)"

def body(first, opening, offer, title):
    q = f"Would you be interested in hearing more about the feature and the option I believe would suit {title} best?"
    return "\n\n".join([f"Hi {first},", opening, SPOT, SERV, offer, q, SIGN])

rows = []
def add(**k):
    r = {c: "" for c in COLS}
    r.update(k)
    r["country"] = "Malaysia"
    r["send_status"] = ""
    r["date_sent"] = ""
    r["has_amazon_a_plus"] = r.get("has_amazon_a_plus") or A_PLUS_NOTE
    if r.get("personalized_opening"):
        r["email_body"] = body(r["first_name"], r["personalized_opening"], k["_offer"], r["book_title"])
    r.pop("_offer", None)
    rows.append(r)

# ---------------- READY ----------------
add(author_name="Malachi Edwin Vethamani", email="mevethamani@gmail.com", first_name="Malachi",
 book_title="Have I Got Something To Tell You",
 book_topic="Short story collection (Penguin Random House SEA, 2024) set from post-independence Malaysia to the Covid-19 years; love, loss, family, sexuality, identity, mixed-race couples, gay men misunderstood by their families.",
 personalization_detail="Collection spans post-independence Malaysia to the Covid-19 years and centres marginalised identities (SARE review). Born in Brickfields, KL; Emeritus Professor, University of Nottingham Malaysia. Launched in KL on 28 April 2024.",
 personalization_reason="The book's historical sweep and focus on overlooked lives is its most distinctive verified feature and gives a sincere, specific opener.",
 personalized_opening="I came across Have I Got Something To Tell You while reading about Malaysian short fiction in English, and I was struck by how the stories travel from post-independence Malaysia all the way to the Covid-19 years. Giving voice to people whose lives are so often left out, from mixed-race couples to gay men misunderstood by their families, feels like deeply human work.",
 _offer="I also noticed that the collection does not currently appear to have an audio edition. Short stories can work beautifully in audio, since a listener can enjoy one complete story in a single sitting, and we can handle the complete production process for you. This is entirely optional and separate from the Spotlight.",
 subject_line="An Author Spotlight for Have I Got Something To Tell You",
 research_sources="https://en.wikipedia.org/wiki/Malachi_Edwin_Vethamani; https://www.nottingham.edu.my/Social-Sciences/People/malachi.vethamani; https://www.malachiedwinvethamani.com/; https://www.malachiedwinvethamani.com/creative-writing-courses/; https://www.malachiedwinvethamani.com/events/; https://www.thepoetrylighthouse.com/poets3/malachi-edwin-vethamani; https://queersoutheastasia.com/queer-southeast-asias-directory-of-writers-and-artists; https://www.penguin.sg/book/have-i-got-something-to-tell-you/; https://www.penguin.com.au/books/have-i-got-something-to-tell-you-9789815144857; http://mjes.um.edu.my/index.php/SARE/article/view/52710; https://www.thestar.com.my/lifestyle/culture/2024/04/25/malaysian-author-highlights-story-sharing-for-stronger-bonds",
 research_confidence="High", personalization_status="Ready",
 has_audiobook="No",
 audiobook_check_sources="Two independent web searches for title/author + audiobook, Audible, Storytel returned only print listings; " + AMZ + "; Audible, Apple Books, Spotify pages could not be opened directly (network policy)",
 has_author_website="Yes", website_url="https://www.malachiedwinvethamani.com/",
 website_quality="Working site with poetry, short stories, creative non-fiction, events, courses, bibliography of Malaysian literature in English, and academic writing.",
 has_book_trailer="No", book_trailer_url="",
 best_offer="Audiobook Production",
 offer_reason="Two separate searches found no audio edition of the 2024 collection; he already has a working website, so audio is the clearest verified gap.")

add(author_name="Tunku Halim (Tunku Halim bin Tunku Abdullah)", email="tunkuhalim@gmail.com", first_name="Tunku Halim",
 book_title="Vermillion Eye",
 book_topic="Gothic horror (Penguin Random House SEA, Nov 2025; first published 2000): a 1957 murder in Kuala Lumpur linked to a man-eating swarm of flies in present-day Sydney; updated for the dark net and influencer era.",
 personalization_detail="New 25th anniversary edition, Nov 2025, updated for the dark net and social media influencers; novel is a study text at the National University of Singapore. Debut Dark Demon Rising nominated for the 1999 International IMPAC Dublin Literary Award.",
 personalization_reason="A freshly reissued anniversary edition is timely and verified, and the premise is strongly visual, which ties naturally to the trailer offer.",
 personalized_opening="I saw that Penguin Random House SEA brought out a 25th anniversary edition of Vermillion Eye last November, updated for the age of the dark net and social media influencers. A story that links a 1957 murder in Kuala Lumpur to a swarm of flies in present-day Sydney is a wonderfully unsettling premise, and it is lovely to see the novel studied at the National University of Singapore.",
 _offer="With the new edition out, a short cinematic trailer and a set of promotional reels could help introduce Vermillion Eye to readers who discover books on TikTok, Instagram, and YouTube Shorts. I was not able to find an official trailer for this edition. These services are optional and completely separate from the Spotlight.",
 subject_line="Featuring the story behind Vermillion Eye",
 research_sources="https://en.wikipedia.org/wiki/Tunku_Halim; https://tunkuhalim.wordpress.com/about/; http://chiasingloh0817.blogspot.com/2014/02/tunku-halim_3697.html; https://www.goodreads.com/author/show/1168589.Tunku_Halim; https://www.smashwords.com/profile/view/tunkuhalim; https://asianreviewofbooks.com/vermillion-eye-by-tunku-halim/; https://www.penguin.sg/book_author/tunku-halim/; https://www.audible.in/author/Tunku-Halim/B001JO8YZY; https://www.youtube.com/watch?v=LOwAkWDKoSg",
 research_confidence="Medium", personalization_status="Ready",
 has_audiobook="Unclear",
 audiobook_check_sources="Audible author page exists with other titles (https://www.audible.in/author/Tunku-Halim/B001JO8YZY); YouTube audiobook preview for Scream to the Shadows; searches found no Vermillion Eye audio edition; " + AMZ,
 has_author_website="Yes", website_url="http://www.tunkuhalim.com/",
 website_quality="Main site plus active blog 'Write Lah!' (tunkuhalim.wordpress.com) with writing tips and contact details on About page. Main site content not directly viewable here.",
 has_book_trailer="No", book_trailer_url="",
 best_offer="Book Trailer and Promotional Reels",
 offer_reason="Two separate searches found no trailer for the Nov 2025 anniversary edition; he already has audio titles and a website, and a new edition of a visual horror premise suits short video.")

add(author_name="Ivy Ngeow", email="ivy_ngeow@yahoo.com", first_name="Ivy",
 book_title="In Safe Hands",
 book_topic="Psychological thriller set in London (Penguin Random House SEA, 21 Oct 2025): Genevieve Ho moves in with her wealthy father after her marriage and business collapse and finds his live-in carer Stella already there. Shortlisted for the Joffe Books Prize 2024.",
 personalization_detail="Born and raised in Johor Bahru; In Safe Hands shortlisted for the Joffe Books Prize 2024; The American Boyfriend has an audiobook; Frost Magazine interview Nov 2025.",
 personalization_reason="The novel's central domestic setup is verified and specific, and the prize shortlisting is a genuine, recent achievement.",
 personalized_opening="I read about In Safe Hands, your psychological thriller from Penguin Random House SEA, and the setup of Genevieve Ho moving back in with her wealthy father, only to find his live-in carer Stella already there, is such a gripping source of tension. It was also great to learn the book was shortlisted for the Joffe Books Prize in 2024.",
 _offer="I noticed that The American Boyfriend has an audiobook edition, but I was not able to find one for In Safe Hands yet. A tense domestic thriller like this can be especially absorbing in audio, and we can manage the complete production process if you would like to explore it. This is entirely optional and separate from the Spotlight.",
 subject_line="An Author Spotlight for In Safe Hands",
 research_sources="https://writengeow.com/; https://writengeow.com/contact/; https://www.penguin.sg/book_author/ivy-ngeow/; https://www.penguin.sg/book/in-safe-hands/; https://www.goodreads.com/en/book/show/234509578-in-safe-hands; https://www.frostmagazine.com/2025/11/interview-with-ivy-ngeow-author-of-in-safe-hands/; https://play.google.com/store/audiobooks/details/The_American_Boyfriend?id=AQAAAEAyyQ2JLM&hl=en_US; https://www.barnesandnoble.com/w/the-american-boyfriend-ivy-ngeow/1143668091?ean=2940190988154; https://www.goodreads.com/videos/169706-overboard-a-trailer",
 research_confidence="Medium", personalization_status="Ready",
 has_audiobook="Unclear",
 audiobook_check_sources="Search for 'In Safe Hands' + audiobook returned only print listings; Google Play and Barnes & Noble show an audiobook for a different title (The American Boyfriend); " + AMZ,
 has_author_website="Yes", website_url="https://writengeow.com/",
 website_quality="Working site with Books, Bio, Contact, Press FAQs, free-book signup, guest posts and interviews.",
 has_book_trailer="No", book_trailer_url="",
 best_offer="Audiobook Production",
 offer_reason="No audio edition of In Safe Hands found, while her previous novel has one; she has a working website and has used trailers before, so audio is the clearest gap. Note: she appears to live in the UK; Malaysian identity based on being born and raised in Johor Bahru. Email is published on her site in 'ivy_ngeow at yahoo dot com' form.")

add(author_name="Anna Tan", email="anna@annatsp.com", first_name="Anna",
 book_title="Amok",
 book_topic="YA fantasy set in a magical Nusantara, following Putera Mikal and the Sultanate of Terang; Book 1 of the Absolution series (Teaspoon Publishing, ISBN 9789671963418).",
 personalization_detail="Penang-based author and editor; President of the Malaysian Writers Society; Chevening-funded MA in Creative Writing; runs indie press Teaspoon Publishing; next book 'Jasmine & the Perfect Brew' due Fall 2026.",
 personalization_reason="Her community role and the book's Nusantara setting are both verified and let the opening feel personal without claiming to have read the book.",
 personalized_opening="I have been reading about Amok and its magical Nusantara setting, following Putera Mikal and the Sultanate of Terang, and I love seeing a YA fantasy built around a world so close to home. Your work as President of the Malaysian Writers Society and through Teaspoon Publishing also shows how much you give back to other Malaysian writers.",
 _offer="Because Amok is set in such a vivid world, I think a short cinematic trailer and a few promotional reels could be a lovely way to introduce it to readers on TikTok, Instagram, and YouTube Shorts. We can take care of the whole process for you. These services are optional and completely separate from the Spotlight.",
 subject_line="Amok and your author journey",
 research_sources="https://www.annatsp.com/; https://www.annatsp.com/contact; https://www.facebook.com/annatansp/; https://malaysianwriterssociety.org/personnel/anna-tan/; https://teaspoonpublishing.com.my/shop/absolution-series/; https://ubsmebooks.com/book/info/341324/Amok; https://www.netgalley.com/catalog/book/303759; https://www.goodreads.com/author/show/6536500.Anna_Tan",
 research_confidence="Medium", personalization_status="Ready",
 has_audiobook="Unclear",
 audiobook_check_sources="Audiobook search not completed (search limit reached); " + AMZ,
 has_author_website="Yes", website_url="https://www.annatsp.com/",
 website_quality="Working site with contact page, mailing list, shop, and edited anthologies; separate blog at blog.annatsp.com.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Book Trailer and Promotional Reels",
 offer_reason="Website is verified, so that is not a gap. Trailer and audio status were not fully checked; the email offers video without claiming a trailer is missing. Confirm no trailer exists before sending if preferred. Publication year not confirmed.")

add(author_name="Anita Harris Satkunananthan", email="aharris@ukm.edu.my", first_name="Anita",
 book_title="Watermyth",
 book_topic="Mythic Gothic fantasy, Book 1 of the Cantata of the Fourfold Realms (Watermaidens Press, 7 Jan 2024): Regya sits the Mermaid Storytelling Exam on an island of exiles and uncovers the cause of a war threatening the merkingdoms.",
 personalization_detail="Debut novel longlisted for the BSFA Award for Best Novel; Senior Lecturer in Literatures in English at UKM; 2016 Rhysling Award nominee; published in Clarkesworld and Strange Horizons (formerly as Nin Harris).",
 personalization_reason="The BSFA longlisting and the storytelling-exam premise are verified and distinctive, and the storytelling theme links naturally to audio.",
 personalized_opening="I was delighted to read about Watermyth, your debut novel and the first book in the Cantata of the Fourfold Realms, and to see it longlisted for the BSFA Award for Best Novel. The idea of Regya uncovering the cause of a war through the stories of the island's watermaidens, all while sitting a Mermaid Storytelling Exam, is such an original premise.",
 _offer="I also noticed that Watermyth does not appear to have an audio edition at the moment. A novel so centred on storytelling feels like a natural fit for listeners, and we can take care of the complete production process if that is something you would like to explore. This is optional and entirely separate from the Spotlight.",
 subject_line="An Author Spotlight for Watermyth",
 research_sources="https://www.ukm.my/fssk/v2/expertise/dr-anita-harris-satkunananthan/; https://www.ukm.my/fssk/v2/?position=senior-lecturer; https://mythopoetica.com/; https://mythopoetica.com/biography/; https://mythopoetica.com/2024/09/01/places-to-purchase-watermyth-ebook/; https://www.goodreads.com/author/show/47791093.Anita_Harris_Satkunananthan; https://madeofstardustandstubbornness.com/2024/09/24/watermyth-review/; https://www.barnesandnoble.com/w/watermyth-anita-harris-satkunananthan/1144573853; https://books.google.com/books/about/Watermyth.html?id=okJx0AEACAAJ",
 research_confidence="Medium", personalization_status="Ready",
 has_audiobook="Unclear",
 audiobook_check_sources="Search for Watermyth + audiobook/trailer returned only ebook and paperback listings (Kobo, Amazon, Indigo); author's buy-links page lists ebook/print; " + AMZ,
 has_author_website="Yes", website_url="https://mythopoetica.com/",
 website_quality="Working author site with biography, bibliography and purchase links for Watermyth.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Audiobook Production",
 offer_reason="Only ebook and print editions surfaced in searches and on her own buy-links page; website exists. Email is her UKM staff address. Based in Malaysia (UKM); citizenship not stated in sources.")

add(author_name="Rajen Devadason", email="rajen@RajenDevadason.com", first_name="Rajen",
 book_title="Financial Freedom 2: Through Malaysian Equities and Unit Trusts",
 book_topic="Personal finance and investing through Malaysian equities and unit trusts (2000); part of the Financial Freedom series, which raised more than RM800,000 for Malaysian charities.",
 personalization_detail="Senior Malaysian financial planner, author and speaker; the Financial Freedom series raised more than RM800,000 for Malaysian charities; NST columnist; CEO of RD WealthCreation and RD Book Projects.",
 personalization_reason="The charitable impact of the series is a verified, meaningful detail that honours his work.",
 personalized_opening="I was reading about your Financial Freedom series and was moved to learn that the books helped raise more than RM800,000 for Malaysian charities. Pairing practical guidance on Malaysian equities and unit trusts with that kind of giving is rare, and your years as a financial planner, columnist, and speaker clearly shape how you write about money.",
 _offer="Since Financial Freedom 2 first appeared in 2000, I think our Book Relaunch Package could be a good fit, with refreshed visuals and promotional materials to introduce your work to a new generation of Malaysian readers. It is completely optional and separate from the Spotlight.",
 subject_line="A feature invitation for Rajen",
 research_sources="http://www.rajendevadason.com/; http://www.rajendevadason.com/about/; http://www.rajendevadason.com/contact; https://learn.rajendevadason.com/; https://asianbeacon.org/resonate-or-isolate/; https://asianbeacon.org/2021/04/16/freedom-joy-and-liberty/; https://www.goodreads.com/author/show/5583983.Rajen_Devadason; https://booksnbobs.com/product/financial-freedom-2-through-malaysian-equities-and-unit-trusts-3/; https://www.nst.com.my/authors/rajen-devadason; https://smartfinance.my/planners/rajen-devadason-cfp",
 research_confidence="Medium", personalization_status="Ready",
 has_audiobook="Unclear",
 audiobook_check_sources="No audiobook evidence surfaced; dedicated audiobook search not completed (search limit reached); " + AMZ,
 has_author_website="Yes", website_url="http://www.rajendevadason.com/",
 website_quality="Working site with About, Contact, free articles and e-newsletter; separate learning site.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Book Relaunch Package",
 offer_reason="Book verified as published in 2000 and no newer title was found, so renewed visibility is the clearest verified gap. Publisher not confirmed.")

# ---------------- NEEDS REVIEW ----------------
add(author_name="Karina Robles Bahrin", email="karina.ts.bahrin@gmail.com", first_name="Karina",
 book_title="The Accidental Malay",
 book_topic="Novel (Epigram Books; Picador/Pan Macmillan edition): Jasmine Leong, heir to a bak kwa company, learns after her grandmother's death that she is partly Malay Muslim; examines the human cost of Malaysia's racial policies. Won the Epigram Books Fiction Prize 2022.",
 personalization_detail="Won the Epigram Books Fiction Prize; lives on Langkawi, where she runs La Pari-Pari and fatCUPID with her sister; co-founded Suatukula community arts initiative; about 20 years in corporate communications.",
 personalization_reason="The prize and her Langkawi life are well documented and give a warm, specific opening.",
 personalized_opening="I loved reading about The Accidental Malay and how Jasmine Leong, heir to a bak kwa empire, discovers she is partly Malay Muslim after her grandmother's death. It was wonderful to see the novel win the Epigram Books Fiction Prize, and I was also interested to learn about your life on Langkawi and your work with Suatukula.",
 _offer="I noticed you currently share your links through Linktree, and a dedicated author website could bring your books, biography, interviews, media information, and purchase links together in one place. We can also prepare a downloadable press kit for events and media requests. These services are optional and entirely separate from the Spotlight.",
 subject_line="An Author Spotlight for The Accidental Malay",
 research_sources="https://en.wikipedia.org/wiki/The_Accidental_Malay; https://linktr.ee/Karinaroblesbahrin; https://www.tatlerasia.com/people/karina-robles-bahrin; https://www.jpf.go.jp/e/project/intel/exchange/yomu/2025/04-02.html; https://epigrambookshop.sg/products/the-accidental-malay; https://www.panmacmillan.com/authors/karina-robles-bahrin/45635; https://www.audible.com/pd/The-Accidental-Malay-Audiobook/B0CYM8KT1Y; https://www.optionstheedge.com/topic/culture/accidental-malay-karina-robles-bahrin%E2%80%99s-award-winning-first-novel-examines-human-cost; https://www.malaymail.com/news/life/2022/09/02/on-menopause-and-being-malay-author-karina-robles-bahrin-opens-up-on-award-winning-debut-novel-the-accidental-malay/26088",
 research_confidence="Medium", personalization_status="Needs Review",
 has_audiobook="Yes",
 audiobook_check_sources="Audible US: https://www.audible.com/pd/The-Accidental-Malay-Audiobook/B0CYM8KT1Y (narrated by Vera Chok, Picador); Audible AU: https://www.audible.com.au/pd/The-Accidental-Malay-Audiobook/B0DM26SSDY",
 has_author_website="No", website_url="",
 website_quality="No official author website found; uses a Linktree link hub (https://linktr.ee/Karinaroblesbahrin).",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Author Website and Press Kit",
 offer_reason="Audiobook already exists; only a Linktree was found, so a central website and press kit is the clearest gap. NEEDS REVIEW: email appeared in search excerpts alongside her Linktree/bio but the exact source page could not be isolated; confirm on the Linktree before sending.")

add(author_name="KC Lau", email="support@kclau.com", first_name="KC",
 book_title="Millionaire Roadmap: Your Step-by-Step Plan to Reach RM1 Million & Beyond",
 book_topic="Personal finance (Ace Premier, Jan 2026): building wealth from zero to RM1 million using the SIR Framework (save, earn, invest), four phases of wealth accumulation, and common traps for Malaysians.",
 personalization_detail="Trained aeronautical engineer turned Registered Financial Planner; described as one of Malaysia's most published personal finance authors; two decades teaching Malaysians; new book published Jan 2026.",
 personalization_reason="The new 2026 book and his SIR Framework are verified and current.",
 personalized_opening="I saw that your new book, Millionaire Roadmap, came out at the start of this year, and I liked how it turns your SIR Framework of saving, earning, and investing into clear phases for Malaysians building wealth. After two decades of teaching personal finance, it feels like a very natural next chapter.",
 _offer="I was not able to find an audio edition of Millionaire Roadmap. Practical finance books can suit listeners who like to learn on the go, and we can manage the complete production process if you would like to offer one. This is optional and completely separate from the Spotlight.",
 subject_line="An Author Spotlight for Millionaire Roadmap",
 research_sources="https://kclau.com/; https://kclau.com/books/; https://kclau.com/webinar/faq/; https://kclau.com/book-reviews/financial-planning-books-magazine-malaysia/; https://www.popularonline.com.my/default/catalog/product/view/_ignore_category/1/id/238841/s/9786297809366/?did=5897; https://mphonline.com/products/millionaire-roadmap-your-step-by-step-plan-to-reach-rm1-million-beyond; https://www.goodreads.com/author/show/20269148.KC_Lau; https://www.youtube.com/playlist?list=PLIpMwVNPtW3rzEXDmsD1Qi8OkqlYcFMtT",
 research_confidence="Medium", personalization_status="Needs Review",
 has_audiobook="Unclear",
 audiobook_check_sources="One search (Millionaire Roadmap + Spotify/Audible/Storytel/Google Play) returned only generic platform pages; " + AMZ,
 has_author_website="Yes", website_url="https://kclau.com/",
 website_quality="Extensive working site with blog, books page, webinars, courses and free ebooks.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Audiobook Production",
 offer_reason="Strong website and YouTube presence, so audio is the most plausible gap, but only one audio search was run. NEEDS REVIEW: the only verified email is a general support inbox (refunds/accounts), not a personal or press address.")

add(author_name="Mawar Safei", email="mawar.safei@ukm.edu.my", first_name="Mawar",
 book_title="Narasi Gua dan Raqim",
 book_topic="Malay-language short story collection (ITBM, 2014 per NTU ACWP listing). Specific theme not confirmed.",
 personalization_detail="SEA Write Award 2018 (Malaysia); 2018 National Academic Award (Art and Creativity); multiple Malaysian literary prizes; Associate Professor, UKM; July 2026 Dewan Sastera piece on how her fiction conveys knowledge implicitly.",
 personalization_reason="The SEA Write Award and the recent Dewan Sastera discussion are verified and recent; the book's theme could not be confirmed so the opening avoids it.",
 personalized_opening="I was delighted to learn that you received the SEA Write Award in 2018, alongside many other honours for your short fiction. It was also lovely to see a recent Dewan Sastera piece discussing how your stories carry knowledge quietly beneath the surface, which made me want to learn more about collections such as Narasi Gua dan Raqim.",
 _offer="I was not able to find audio editions of your short story collections. Short fiction can reach new listeners through audio, and we can handle the complete production process if that is something you would like to explore. This is optional and entirely separate from the Spotlight.",
 subject_line="A feature invitation for Mawar",
 research_sources="https://www.ukm.my/fssk/v2/?staff=prof-madya-dr-mawar-safei; https://blogs.ntu.edu.sg/acwp/mawar-safei/; https://www.goodreads.com/author/show/4579505.Mawar_Safei; https://www.researchgate.net/profile/Mawar-Safei; http://mawarshafei.blogspot.com/; http://mawarshafei.blogspot.com/p/antologikumpulan-cerpen.html; https://dewansastera.jendeladbp.my/2026/07/15/22024/",
 research_confidence="Medium", personalization_status="Needs Review",
 has_audiobook="Unclear",
 audiobook_check_sources="Search for 'Mawar Safei buku audio OR audiobook' returned only generic channels; " + AMZ,
 has_author_website="Poor", website_url="http://mawarshafei.blogspot.com/",
 website_quality="Personal Blogspot blog with a list of anthologies and collections; not a dedicated author website.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Audiobook Production",
 offer_reason="No audio editions found. NEEDS REVIEW: she writes in Malay (consider sending in Malay), the book's theme is unconfirmed, and she was born in Singapore though she works at UKM and won the SEA Write Award representing Malaysia. A website offer is a reasonable alternative since she only has a blog.")

add(author_name="Emila Yusof", email="emilayusof@gmail.com", first_name="Emila",
 book_title="My Mother's Garden",
 book_topic="Children's picture book (Oyez!Books, with Daphne Lee) about a girl who delights in her mother's garden of flowers, bees, dragonflies and butterflies; published in English, Malay, Chinese and German.",
 personalization_detail="Kuala Lumpur-based author-illustrator with over 100 published titles; Pohon Kenangan selected for White Ravens 2025; nominated for the Astrid Lindgren Memorial Award.",
 personalization_reason="The book's multilingual reach and her White Ravens selection are verified and celebratory.",
 personalized_opening="I was delighted to learn how My Mother's Garden, with its flowers, bees, dragonflies, and butterflies, has travelled into English, Malay, Chinese, and German editions. It was also wonderful to read that Pohon Kenangan was selected for the White Ravens 2025 list, and that you have now published more than 100 titles.",
 _offer="Picture books are so visual that I think a short book trailer and a few promotional reels could be a lovely way to share your stories with parents and teachers on Instagram, TikTok, and YouTube Shorts. I was not able to find an official trailer for My Mother's Garden, and we would be happy to handle the whole process. These services are optional and entirely separate from the Spotlight.",
 subject_line="Featuring the story behind My Mother's Garden",
 research_sources="https://emilayusof.com/; https://emilayusof.com/category/events/awareness/; https://emilayusof.com/author/emila/page/298/?wpmp_switcher=desktop; https://mirrorswindowsdoors.org/wp/interview-emila-yusof/; https://www.freemalaysiatoday.com/category/leisure/2024/08/30/book-illustrator-crafts-malaysian-stories-for-children; https://www.goodreads.com/book/show/8820944-my-mother-s-garden",
 research_confidence="Low", personalization_status="Needs Review",
 has_audiobook="Unclear",
 audiobook_check_sources="No audiobook found in searches; book exists in animation format; " + AMZ,
 has_author_website="Yes", website_url="https://emilayusof.com/",
 website_quality="Blog 'Emilatopia' with FAQ, Books and Experience pages.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Book Trailer and Promotional Reels",
 offer_reason="Picture books are highly visual and no official trailer was found. NEEDS REVIEW: email comes only from old blog posts and may be outdated; publication year of the book not confirmed; animation versions exist, so confirm a trailer is genuinely missing.")

add(author_name="Golda Mowe", email="", first_name="Golda",
 book_title="The Monk Prince",
 book_topic="Historical fantasy (Penguin Random House SEA, Nov 2021) set in 8th-century Santubong, Borneo: a prince raised by Buddhist monks is pushed to give up non-violence.",
 personalization_detail="Born and raised in Sarawak to an Iban mother and Melanau father; lives in Sibu; Iban Dream trilogy draws on real Iban beliefs; Waseda University graduate who left corporate work in 2004 to write.",
 personalization_reason="Her Sarawak heritage and the book's Borneo setting are verified and closely connected.",
 personalized_opening="I was fascinated to read about The Monk Prince and its setting in 8th-century Santubong, where a prince raised by Buddhist monks is pushed to give up his commitment to non-violence. Knowing that you grew up in Sarawak with Iban and Melanau heritage, and that your Iban Dream trilogy draws on real Iban beliefs, makes your work feel especially rooted in Borneo.",
 _offer="I was not able to find an audio edition of The Monk Prince. A sweeping historical fantasy like this could be wonderful to listen to, and we can handle the complete production process if you would like to explore it. This is optional and separate from the Spotlight.",
 subject_line="An Author Spotlight for The Monk Prince",
 research_sources="https://www.penguin.sg/book_author/golda-mowe/; https://www.penguin.sg/book/the-monk-prince/; https://www.riwayat.my/penguin/golda-mowe-the-monk-prince-fantasy-penguin; https://www.goodreads.com/author/show/6460446.Golda_Mowe; https://rimbundahan.org/golda-mowe/; http://borneoexpatwriter.blogspot.com/2022/02/sarawak-author-golda-mowe-featured-in.html; https://www.overdrive.com/media/1081789/iban-dream",
 research_confidence="Low", personalization_status="Needs Review",
 has_audiobook="No",
 audiobook_check_sources="Searches found no audio edition of The Monk Prince; one search excerpt suggested Iban Dream is on Audible (unconfirmed); " + AMZ,
 has_author_website="Unclear", website_url="",
 website_quality="A site (gmowe.ws) and blog (tipsbygolda.blogspot.com) were mentioned in search excerpts but could not be verified.",
 has_book_trailer="No", book_trailer_url="",
 best_offer="Audiobook Production",
 offer_reason="No audio edition of The Monk Prince found. NEEDS REVIEW: no verified public email. An address appeared in a single search excerpt but did not verify on an exact-match search, so it is not recorded.")

add(author_name="Cassandra Aasmundsen-Fry", email="", first_name="Cassandra",
 book_title="The Glory in Us All: Embracing the Life We Bury",
 book_topic="Psychology/non-fiction (Penguin Random House SEA, Feb 2025): intergenerational trauma, legacies and resilience, told through clients' stories.",
 personalization_detail="US-trained clinical psychologist practising in Kuala Lumpur; founded MindWell in Mont Kiara; book featured in Free Malaysia Today's '4 illuminating Malaysian non-fiction books in 2025'.",
 personalization_reason="The FMT feature and her KL practice are verified and directly tied to the book.",
 personalized_opening="I came across The Glory in Us All in Free Malaysia Today's list of illuminating Malaysian non-fiction books for 2025, and I was drawn to how it explores intergenerational trauma and resilience through the stories of people you have worked with as a clinical psychologist in Kuala Lumpur.",
 _offer="I noticed that your online presence is centred on your MindWell practice, and a dedicated author website could bring your book, biography, interviews, media information, and purchase links together in one place for readers and the press. We can also prepare a downloadable press kit. These services are optional and entirely separate from the Spotlight.",
 subject_line="Featuring the story behind The Glory in Us All",
 research_sources="https://www.penguin.sg/book/the-glory-in-us-all/; https://www.freemalaysiatoday.com/category/leisure/2025/01/18/4-illuminating-malaysian-non-fiction-books-in-2025; https://www.mindwell.com.my/about-me/; https://www.mindwell.com.my/",
 research_confidence="Low", personalization_status="Needs Review",
 has_audiobook="Unclear",
 audiobook_check_sources="Not checked (search limit reached); " + AMZ,
 has_author_website="Unclear", website_url="",
 website_quality="Practice website https://www.mindwell.com.my/ (About Me, blog); no separate author site found in limited checks.",
 has_book_trailer="Unclear", book_trailer_url="",
 best_offer="Author Website and Press Kit",
 offer_reason="Only a clinical practice site was found. NEEDS REVIEW: no public email found (only masked addresses on data-broker sites); subtitle varies between listings ('Embracing the Life We Bury' vs 'Embracing the Life and Knowledge We Bury'); based in Malaysia, not Malaysian-born.")

add(author_name="Faisal Tehrani (Mohd Faizal Musa)", email="", first_name="Faizal",
 book_title="The Koro Riots",
 book_topic="English translation published by Penguin Random House SEA, 2023 (per Wikipedia). Theme not confirmed in research.",
 personalization_detail="Born in Kuala Lumpur, 1974; National Art Award 2006; Research Fellow at ATMA, UKM; Associate Fellow at Harvard's Weatherhead Center.",
 personalization_reason="",
 research_sources="https://en.wikipedia.org/wiki/Faisal_Tehrani; https://www.ukm.my/atma/en/expertise/dr-mohd-faizal-musa-faisal-tehrani/; https://www.facebook.com/gerakbudaya/posts/the-professor-by-mohd-faizal-musa-faisal-tehrani-now-in-english-translated-by-br/10156745967484426/",
 research_confidence="Low", personalization_status="Needs Review",
 has_audiobook="Unclear", audiobook_check_sources="Not checked (search limit reached)",
 has_author_website="Unclear", website_quality="Not checked.",
 has_book_trailer="Unclear",
 best_offer="", offer_reason="No personal email found (only ATMA institute general inboxes). Book details rely on Wikipedia only. No email drafted until email and book theme are verified.")

# ---------------- EXCLUDED ----------------
add(author_name="Samantha Chong", email="", first_name="Samantha",
 book_title="Prodigal Tiger",
 book_topic="Malaysian YA fantasy (Penguin/PRH, 17 March 2026).",
 personalization_detail="", personalization_reason="",
 research_sources="https://www.hisamchong.com/; https://www.hisamchong.com/contact-samantha-chong; https://www.audible.com/pd/Prodigal-Tiger-Audiobook/B0FDXDNF9C",
 research_confidence="Medium", personalization_status="Excluded",
 has_audiobook="Yes", audiobook_check_sources="Audible: https://www.audible.com/pd/Prodigal-Tiger-Audiobook/B0FDXDNF9C",
 has_author_website="Yes", website_url="https://www.hisamchong.com/", website_quality="Working author site with contact form; publishing enquiries routed to her literary agent.",
 has_book_trailer="Unclear",
 best_offer="", offer_reason="Excluded: no public email (contact form only; enquiries directed to agent), US-based, and already has an audiobook with a major publisher.")

add(author_name="Joyce Nga Koe Hwee", email="", first_name="Joyce",
 book_title="Personal Financial Planning: A Comprehensive Guide to Personal Financial Planning in Malaysia",
 book_topic="Personal financial planning guide (Sunway University Press; year not found).",
 personalization_detail="", personalization_reason="",
 research_sources="https://sunwayuniversity.edu.my/sunway-business-school/staff-profiles/dr-joyce-nga-koe-hwee; https://sunwayuniversity.edu.my/press/books/personal-financial-planning",
 research_confidence="Low", personalization_status="Excluded",
 has_audiobook="Unclear", audiobook_check_sources="Not checked",
 has_author_website="Unclear", has_book_trailer="Unclear",
 best_offer="", offer_reason="Excluded: email obfuscated on staff profile and not recoverable; book is an academic/professional guide rather than a general-audience title.")

# ---------------- checks ----------------
for r in rows:
    for c in COLS:
        assert "—" not in r[c] and "–" not in r[c], (r["author_name"], c)
    if r["email_body"]:
        low = r["email_body"].lower()
        for bad in ["ai ", "artificial", "synthetic", "automated", "voice clon", "text-to-speech", "generative", "guarantee"]:
            assert bad not in low, (r["author_name"], bad)
        n = len(re.findall(r"\b[\w'&:.-]+\b", r["email_body"]))
        print(f'{r["personalization_status"]:13} {n:4} words  {r["author_name"]}')
    else:
        print(f'{r["personalization_status"]:13}    - no body  {r["author_name"]}')

with open(f"{OUT}/malaysia_author_outreach.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=COLS, quoting=csv.QUOTE_ALL)
    w.writeheader(); w.writerows(rows)

wb = Workbook(); ws = wb.active; ws.title = "Malaysia Authors"
ws.append(COLS)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
fills = {"Ready": "E2F0D9", "Needs Review": "FFF2CC", "Excluded": "F4CCCC"}
for r in rows:
    ws.append([r[c] for c in COLS])
    st = ws.cell(ws.max_row, COLS.index("personalization_status") + 1)
    st.fill = PatternFill("solid", fgColor=fills[r["personalization_status"]])
widths = {"email_body": 80, "personalized_opening": 60, "research_sources": 60, "book_topic": 45,
          "personalization_detail": 45, "offer_reason": 50, "audiobook_check_sources": 45, "website_quality": 40}
for i, c in enumerate(COLS, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(c, 22)
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"
wb.save(f"{OUT}/malaysia_author_outreach.xlsx")
print("rows:", len(rows))
