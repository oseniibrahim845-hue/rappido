"""Merge segment JSONL files, clean, dedupe, and export the prospect workbook."""
import glob, json, re, os, collections
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Fintech_Trading_Dashboard_Prospects_700.xlsx")
TARGET = 700
COLUMNS = ["Name", "Company", "Website", "Email", "Subject", "Body", "Status", "Country",
           "Industry", "Problem", "Opportunity", "Source", "Evidence", "Notes"]

SENDER = "Oseni Ibrahim\nTrading Bot Developer\nibrahimoseni063@gmail.com"
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@([A-Za-z0-9-]+\.)+[A-Za-z]{2,}$")
FREEMAIL = {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "icloud.com", "proton.me",
            "protonmail.com", "aol.com", "mail.ru", "yandex.ru", "qq.com", "163.com"}
PLACEHOLDER_RE = re.compile(r"\[[^\]]*\]|\{[^}]*\}|<[^>]*>|XXX|TBD", re.I)
SECOND_LEVEL = {"co", "com", "net", "org", "gov", "ac", "ltd", "plc"}


def base_domain(host):
    host = (host or "").lower().strip().rstrip(".")
    if host.startswith("www."):
        host = host[4:]
    parts = host.split(".")
    if len(parts) >= 3 and parts[-2] in SECOND_LEVEL and len(parts[-1]) == 2:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def site_domain(url):
    url = (url or "").strip()
    if not url:
        return ""
    if "://" not in url:
        url = "https://" + url
    return base_domain(urlparse(url).hostname or "")


def norm_company(name):
    n = re.sub(r"[^a-z0-9 ]", " ", (name or "").lower())
    n = re.sub(r"\b(ltd|limited|inc|llc|plc|gmbh|ag|sa|bv|pty|corp|corporation|co|group|"
               r"holdings|technologies|technology|the)\b", " ", n)
    return re.sub(r"\s+", "", n)


def clean(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def sentence(s):
    s = clean(s).rstrip(".")
    return s[0].lower() + s[1:] if s else s


def make_subject(r):
    ind = r["industry"].lower()
    if "crypto" in ind:
        topic = "crypto trading dashboards & automation"
    elif "broker" in ind or "crm" in ind:
        topic = "broker dashboards & trading integrations"
    elif "portfolio" in ind or "wealth" in ind:
        topic = "portfolio & reporting dashboards"
    elif "risk" in ind or "oms" in ind or "ems" in ind or "regtech" in ind:
        topic = "risk monitoring & reporting dashboards"
    elif "prop" in ind:
        topic = "trader monitoring dashboards & automation"
    elif "data" in ind or "api" in ind or "infrastructure" in ind:
        topic = "API integrations & dashboards"
    else:
        topic = "trading dashboards & automation"
    return f"{r['company']}: {topic}"


def make_body(r):
    greet = f"Hi {r['name'].split()[0]}," if r["name"] else f"Hi {r['company']} team,"
    return (
        f"{greet}\n\n"
        f"I came across {r['company']} while researching {r['industry'].lower()} companies "
        f"and noticed that {sentence(r['evidence'])}.\n\n"
        f"Teams in this space often tell me that {sentence(r['problem'])}. "
        f"I build custom trading bots, real-time trading dashboards, risk and portfolio monitoring, "
        f"reporting automation and broker/exchange API integrations, and I think I could help: "
        f"{sentence(r['opportunity'])}.\n\n"
        f"If useful, I can put together a short scoped proposal or a quick prototype so you can see "
        f"the approach before committing to anything. Would you be open to a 15-minute call next week?\n\n"
        f"Best regards,\n{SENDER}"
    )


def main():
    rows, stats = [], collections.Counter()
    for path in sorted(glob.glob(os.path.join(HERE, "segments", "*.jsonl"))):
        for i, line in enumerate(open(path, encoding="utf-8")):
            line = line.strip()
            if not line:
                continue
            try:
                d = json.loads(line)
            except json.JSONDecodeError:
                stats["bad_json_lines"] += 1
                continue
            d = {k: clean(d.get(k, "")) for k in
                 ["company", "website", "country", "industry", "name", "email", "email_source",
                  "source", "evidence", "problem", "opportunity", "notes"]}
            d["segment"] = os.path.basename(path)
            stats["raw_rows"] += 1
            rows.append(d)

    # Required fields: a prospect must have verifiable source + evidence.
    kept = []
    for d in rows:
        if not (d["company"] and d["website"] and d["source"].startswith("http") and d["evidence"]
                and d["problem"] and d["opportunity"]):
            stats["dropped_missing_fields"] += 1
            continue
        if not d["website"].startswith("http"):
            d["website"] = "https://" + d["website"]
        d["domain"] = site_domain(d["website"])
        d["country"] = d["country"] or "Unknown"
        d["industry"] = d["industry"] or "Fintech"
        kept.append(d)

    # Duplicate companies (same domain or same normalized name); keep the row with an email.
    kept.sort(key=lambda d: (0 if d["email"] else 1))
    seen_dom, seen_name, uniq = set(), set(), []
    for d in kept:
        nn = norm_company(d["company"])
        if d["domain"] in seen_dom or nn in seen_name:
            stats["duplicate_companies_removed"] += 1
            continue
        seen_dom.add(d["domain"]); seen_name.add(nn)
        uniq.append(d)

    # Email validation + duplicate email removal.
    seen_email = set()
    for d in uniq:
        e = d["email"].lower().strip().strip(".").replace("mailto:", "")
        d["email"] = ""
        if not e:
            continue
        edom = base_domain(e.split("@")[-1]) if "@" in e else ""
        reason = None
        if not EMAIL_RE.match(e):
            reason = "malformed"
        elif edom in FREEMAIL:
            reason = "free-mail domain"
        elif edom != d["domain"] and edom.split(".")[0] != d["domain"].split(".")[0]:
            reason = f"domain {edom} does not match website"
        if reason:
            stats["invalid_emails_removed"] += 1
            d["notes"] = (d["notes"] + f" | Email removed ({reason}): {e}").strip(" |")
            continue
        if e in seen_email:
            stats["duplicate_emails_removed"] += 1
            d["notes"] = (d["notes"] + f" | Duplicate email removed: {e}").strip(" |")
            continue
        seen_email.add(e)
        d["email"] = e
        if d["email_source"]:
            d["notes"] = (d["notes"] + f" | Email seen at: {d['email_source']}").strip(" |")

    final = uniq[:TARGET]
    stats["excess_rows_trimmed"] = max(0, len(uniq) - TARGET)

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    ws.append(COLUMNS)
    for d in final:
        subj, body = make_subject(d), make_body(d)
        assert not PLACEHOLDER_RE.search(subj + body), (d["company"], body)
        ws.append([d["name"], d["company"], d["website"], d["email"], subj, body, "",
                   d["country"], d["industry"], d["problem"], d["opportunity"], d["source"],
                   d["evidence"], d["notes"]])
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
    widths = [18, 28, 32, 34, 45, 80, 10, 16, 24, 50, 50, 50, 60, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    countries = collections.Counter(d["country"] for d in final)
    industries = collections.Counter(d["industry"] for d in final)
    with_email = sum(1 for d in final if d["email"])
    summary = wb.create_sheet("Summary")
    lines = [("Total prospects", len(final)), ("With verified/public email", with_email),
             ("Without email", len(final) - with_email)]
    lines += [(k.replace("_", " ").capitalize(), v) for k, v in stats.items()]
    for k, v in lines:
        summary.append([k, v])
    summary.append([]); summary.append(["Country", "Prospects"])
    for k, v in countries.most_common():
        summary.append([k, v])
    summary.append([]); summary.append(["Industry", "Prospects"])
    for k, v in industries.most_common():
        summary.append([k, v])
    summary.column_dimensions["A"].width = 40
    wb.save(OUT)

    report = {"total": len(final), "with_email": with_email, "without_email": len(final) - with_email,
              **stats, "countries": dict(countries.most_common()),
              "industries": dict(industries.most_common())}
    json.dump(report, open(os.path.join(HERE, "report.json"), "w"), indent=2)
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
