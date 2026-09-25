import csv, glob, json, os, re, sys
from collections import Counter
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/user/rappido/prospects"
os.makedirs(OUT, exist_ok=True)

COLS = ["Lead ID", "Company Name", "Website", "Country", "City", "Industry", "Lead Category",
        "Company Description", "Contact Name", "Job Title", "Business Email", "Email Source",
        "LinkedIn", "Company LinkedIn", "Evidence URL", "Evidence Type", "Evidence Summary",
        "Likely Business Problem", "Potential Dashboard Solution", "Personalization Angle",
        "Lead Quality", "Email Status", "Notes"]
KEYS = ["company_name", "website", "country", "city", "industry", "lead_category",
        "company_description", "contact_name", "job_title", "business_email", "email_source",
        "linkedin", "company_linkedin", "evidence_url", "evidence_type", "evidence_summary",
        "likely_business_problem", "potential_dashboard_solution", "personalization_angle",
        "lead_quality", "email_status", "notes"]
PROBLEMS = ["Client reporting", "CRM visibility", "Sales pipeline visibility", "Data integration",
            "Manual reporting", "Automation monitoring", "AI agent monitoring", "Ecommerce operations",
            "Order monitoring", "Business reporting", "API monitoring", "Multi client reporting",
            "Data quality", "Internal operations", "Trading analytics", "Other"]
COUNTRIES = {"United States", "Canada", "United Kingdom", "Australia", "Germany", "Netherlands",
             "Switzerland", "Sweden", "Norway", "Denmark", "France", "Ireland", "Belgium", "Austria",
             "Singapore", "New Zealand"}
COUNTRY_ALIASES = {"usa": "United States", "us": "United States", "u.s.": "United States",
                   "united states of america": "United States", "uk": "United Kingdom",
                   "england": "United Kingdom", "scotland": "United Kingdom", "wales": "United Kingdom",
                   "northern ireland": "United Kingdom", "great britain": "United Kingdom",
                   "the netherlands": "Netherlands", "holland": "Netherlands"}
CONSUMER = {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "icloud.com", "aol.com",
            "live.com", "proton.me", "protonmail.com", "gmx.de", "web.de", "yahoo.co.uk", "me.com"}
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")
SPAMMY = re.compile(r"(noreply|no-reply|donotreply|example\.|test@|abuse@|postmaster@|privacy@|gdpr@|dpo@|legal@|unsubscribe)", re.I)


def domain(u):
    if not u:
        return ""
    u = u.strip()
    if not u.startswith("http"):
        u = "https://" + u
    h = urlparse(u).netloc.lower().split(":")[0]
    return h[4:] if h.startswith("www.") else h


def root(d):
    parts = d.split(".")
    if len(parts) >= 3 and parts[-2] in ("co", "com", "org", "net", "ac") and len(parts[-1]) == 2:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def norm_name(n):
    n = re.sub(r"\(.*?\)", "", n.lower())
    n = re.sub(r"[^a-z0-9 ]", " ", n)
    n = re.sub(r"\b(ltd|limited|inc|llc|gmbh|bv|b v|ag|ab|as|aps|pty|plc|sa|sas|co|corp|the|group|agency|digital)\b", " ", n)
    return re.sub(r"\s+", " ", n).strip()


records, rejects = [], []
for f in sorted(glob.glob(os.path.join(SRC, "cat*.json"))):
    try:
        data = json.load(open(f, encoding="utf-8"))
    except Exception as e:
        print("BAD JSON", f, e)
        continue
    for r in data:
        r = {k: (str(r.get(k) or "").strip()) for k in KEYS}
        r["_file"] = os.path.basename(f)
        records.append(r)

seen_dom, seen_name, seen_email = set(), set(), set()
clean = []
for r in records:
    why = []
    d = domain(r["website"])
    if not r["company_name"]:
        why.append("missing company name")
    if not r["evidence_url"].startswith("http"):
        why.append("missing evidence URL")
    if not r["evidence_summary"]:
        why.append("missing evidence summary")
    if why:
        rejects.append((r, "; ".join(why)))
        continue
    if r["_file"] == "cat15.json" and r["country"]:
        r["notes"] = (r["notes"] + " | " if r["notes"] else "") + f"country blanked in QC (was not confirmed by a source; researcher suggested {r['country']})"
        r["country"] = ""
    has_site = bool(d and "." in d and not any(x in d for x in ("workato.com", "linkedin.com", "greenhouse.io", "lever.co", "ashbyhq.com")))
    r["website"] = ("https://" + d) if has_site else ""
    if not has_site:
        r["notes"] = (r["notes"] + " | " if r["notes"] else "") + "company website not confirmed in search results"
    c = COUNTRY_ALIASES.get(r["country"].lower(), r["country"])
    r["country"] = c
    if c and c not in COUNTRIES:
        rejects.append((r, f"country outside target list: {c}"))
        continue
    rd = root(d) if has_site else ""
    nn = norm_name(r["company_name"])
    if (rd and rd in seen_dom) or (nn and nn in seen_name):
        rejects.append((r, "duplicate company/domain"))
        continue
    # email hygiene
    e = r["business_email"].lower().strip().rstrip(".")
    if e:
        ed = e.split("@")[-1]
        bad = None
        if not EMAIL_RE.match(e) or "*" in e:
            bad = "malformed/masked email removed"
        elif SPAMMY.search(e):
            bad = "non-outreach address removed"
        elif ed in CONSUMER:
            bad = "consumer email removed"
        elif e in seen_email:
            bad = "duplicate email removed"
        elif not r["email_source"].startswith("http"):
            bad = "email without source removed"
        if bad:
            r["notes"] = (r["notes"] + " | " if r["notes"] else "") + bad + f" ({e})"
            e = ""
    r["business_email"] = e
    if not e:
        r["email_source"] = ""
        r["email_status"] = "Not found"
    else:
        src_dom = domain(r["email_source"])
        if rd and root(src_dom) == rd:
            r["email_status"] = "Published on company site (via search index)"
        else:
            r["email_status"] = "Listed in third-party directory (verify before sending)"
        if "source page unconfirmed" in r["notes"].lower() or "verify" in r["notes"].lower() and r["company_name"].startswith("Passion Digital"):
            r["email_status"] = "Seen in search index, source page unconfirmed (verify before sending)"
        seen_email.add(e)
    if r["likely_business_problem"] not in PROBLEMS:
        m = [p for p in PROBLEMS if p.lower() == r["likely_business_problem"].lower()]
        r["likely_business_problem"] = m[0] if m else "Other"
    q = r["lead_quality"].capitalize()
    if q not in ("High", "Medium", "Low"):
        q = "Medium"
    # enforce High definition: needs email or named contact
    if q == "High" and not (e or r["contact_name"]):
        q = "Medium"
    if not has_site and q != "Low":
        q = "Low"
    r["lead_quality"] = q
    if not r["contact_name"]:
        r["job_title"] = ""
    for k in ("linkedin", "company_linkedin"):
        if r[k] and "linkedin.com" not in r[k]:
            r[k] = ""
    if rd:
        seen_dom.add(rd)
    if nn:
        seen_name.add(nn)
    clean.append(r)

cat_order = ["HubSpot / CRM Partner", "RevOps / CRM Consulting", "Marketing Agency", "SEO / PPC Agency",
             "Automation Agency", "AI Automation / AI Agents", "B2B SaaS", "Sales Outsourcing / Lead Generation",
             "Shopify / Ecommerce Agency", "Ecommerce Brand / Operator", "Sales Team / Large CRM Pipeline",
             "API / Integration Company", "Trading / Fintech", "BI / Reporting Consultancy",
             "Technology / Operations Company"]
qo = {"High": 0, "Medium": 1, "Low": 2}
clean.sort(key=lambda r: (cat_order.index(r["lead_category"]) if r["lead_category"] in cat_order else 99,
                          qo[r["lead_quality"]], r["company_name"].lower()))
rows = []
for i, r in enumerate(clean, 1):
    rows.append([f"DP-{i:04d}"] + [r[k] for k in KEYS])

base = os.path.join(OUT, "Dashboard_Development_Prospects_700")
with open(base + ".csv", "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(COLS)
    w.writerows(rows)

wb = Workbook()
ws = wb.active
ws.title = "Prospects"
ws.append(COLS)
for row in rows:
    ws.append(row)
hdr = PatternFill("solid", fgColor="1F3864")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = hdr
    c.alignment = Alignment(wrap_text=True, vertical="center")
widths = {"Lead ID": 10, "Company Name": 28, "Website": 30, "Country": 15, "City": 14, "Industry": 22,
          "Lead Category": 26, "Company Description": 50, "Evidence Summary": 60, "Personalization Angle": 55,
          "Evidence URL": 45, "Email Source": 40, "Notes": 40, "Business Email": 30}
for i, col in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 22)
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions
fills = {"High": "C6EFCE", "Medium": "FFEB9C", "Low": "F4CCCC"}
qi = COLS.index("Lead Quality") + 1
for rr in range(2, ws.max_row + 1):
    c = ws.cell(rr, qi)
    c.fill = PatternFill("solid", fgColor=fills.get(c.value, "FFFFFF"))
    for col in ("Website", "Evidence URL", "Email Source", "LinkedIn", "Company LinkedIn"):
        cell = ws.cell(rr, COLS.index(col) + 1)
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")

# summary sheet
s = wb.create_sheet("Summary")
emails = [r["business_email"] for r in clean if r["business_email"]]
stats = {
    "Total prospects": len(clean),
    "Unique companies": len({norm_name(r["company_name"]) for r in clean}),
    "Unique domains": len({root(domain(r["website"])) for r in clean if r["website"]}),
    "Leads without confirmed website": sum(1 for r in clean if not r["website"]),
    "Leads without confirmed country": sum(1 for r in clean if not r["country"]),
    "Unique emails": len(set(emails)),
    "Leads with a public business email": len(emails),
    "  - published on company site (via search index)": sum(1 for r in clean if r["email_status"].startswith("Published")),
    "  - listed in third-party directory": sum(1 for r in clean if r["email_status"].startswith("Listed")),
    "Missing emails": len(clean) - len(emails),
    "Leads with named contact": sum(1 for r in clean if r["contact_name"]),
    "High quality": sum(1 for r in clean if r["lead_quality"] == "High"),
    "Medium quality": sum(1 for r in clean if r["lead_quality"] == "Medium"),
    "Low quality": sum(1 for r in clean if r["lead_quality"] == "Low"),
    "Raw records collected": len(records),
    "Records removed in QC": len(rejects),
}
s.append(["Metric", "Value"])
for k, v in stats.items():
    s.append([k, v])
s.append([])
s.append(["Lead Category", "Count"])
for k, v in Counter(r["lead_category"] for r in clean).most_common():
    s.append([k, v])
s.append([])
s.append(["Country", "Count"])
for k, v in Counter(r["country"] for r in clean).most_common():
    s.append([k, v])
s.column_dimensions["A"].width = 50
s.column_dimensions["B"].width = 12
for row in s.iter_rows():
    if row[0].value in ("Metric", "Lead Category", "Country"):
        for c in row:
            c.font = Font(bold=True)
wb.save(base + ".xlsx")

json.dump({"stats": stats,
           "by_category": Counter(r["lead_category"] for r in clean),
           "by_country": Counter(r["country"] for r in clean),
           "by_quality_category": {f"{a} | {b}": n for (a, b), n in
                                   Counter((r["lead_category"], r["lead_quality"]) for r in clean).items()},
           "rejects": Counter(w for _, w in rejects)},
          open(os.path.join(SRC, "stats.json"), "w"), indent=2)
print(json.dumps(stats, indent=2))
print("by category:", dict(Counter(r["lead_category"] for r in clean)))
print("by country:", dict(Counter(r["country"] for r in clean)))
print("rejects:", dict(Counter(w for _, w in rejects)))
