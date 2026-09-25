"""Build the Apps Script sender spreadsheet from the prospect database.

Usage: python3 raw/export_outreach.py
Reads  Dashboard_Development_Prospects_700.csv and raw/enrichment/*.json
Writes Dashboard_Outreach_Sender.xlsx (+ .csv of the send-ready sheet)
"""
import csv, glob, json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
COLS = ["Name", "Company", "Website", "Email", "Subject", "Body", "Status",
        "Country", "Industry", "Source", "Notes"]
SENDER = ("Oseni Ibrahim", "Trading Bot Developer", "ibrahimoseni063@gmail.com")

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}$")
PLACEHOLDER = re.compile(r"(example\.|@domain\.|@company\.|@email\.|yourname|firstname|lastname|test@|"
                         r"noreply|no-reply|donotreply|privacy@|gdpr@|dpo@|legal@|abuse@|postmaster@|"
                         r"jobs@|careers@|recruit|hr@|\*)", re.I)
CONSUMER = {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "icloud.com", "aol.com", "live.com",
            "proton.me", "protonmail.com", "gmx.de", "web.de", "yahoo.co.uk", "me.com"}
BAD_TEXT = re.compile(r"\[[^\]]*\]|\{\{|\}\}|<[A-Za-z ]+>")

# Problem-specific lines: phrased as common patterns, never as claims about the prospect.
PAIN = {
    "Client reporting": ("Client reporting across several accounts often means pulling numbers from different "
                         "platforms by hand each month.", "cut the time your team spends assembling client reports"),
    "Multi client reporting": ("When you look after many client accounts, getting one clear view across all of "
                               "them is usually harder than it should be.", "give you one view across every client account"),
    "Manual reporting": ("Recurring reports are often still built manually from exports and spreadsheets.",
                         "replace manual exports with numbers that update on their own"),
    "CRM visibility": ("CRM data is only as useful as the reporting on top of it, and native reports rarely show "
                       "everything a team needs.", "give clearer visibility into CRM activity and data quality"),
    "Sales pipeline visibility": ("Pipeline and activity numbers tend to be spread across CRMs, sheets and outreach "
                                  "tools.", "show pipeline, meetings and conversion in one live view"),
    "Data integration": ("Integration work usually leaves teams needing a clear view of what is syncing, what "
                         "failed and why.", "show sync status, errors and data flow across connected systems"),
    "API monitoring": ("Once several APIs and integrations are live, spotting failures early becomes the hard part.",
                       "surface API errors, latency and failed syncs before clients notice"),
    "Automation monitoring": ("As the number of workflows grows, knowing which automations ran, failed or slowed "
                              "down gets difficult.", "track workflow runs, failures and time saved per client"),
    "AI agent monitoring": ("Once AI agents are in production, teams and clients want to see what the agents are "
                            "actually doing and what they cost.", "track conversations, outcomes, errors and token cost for each agent"),
    "Ecommerce operations": ("Store, ad and fulfilment data usually lives in separate tools.",
                             "bring orders, revenue, ad spend and stock into one live view"),
    "Order monitoring": ("Order data often has to be checked across the store, ERP and fulfilment tools.",
                         "show order status, exceptions and fulfilment times in one place"),
    "Business reporting": ("Many BI teams get requests for web-based dashboards or client portals that go beyond "
                           "what standard BI tools do well.", "deliver custom web dashboards and client portals alongside your BI work"),
    "Data quality": ("Reporting is only reliable when the underlying data is clean and complete.",
                     "flag missing, duplicate and stale records automatically"),
    "Internal operations": ("Operations teams often end up switching between many internal systems to answer simple "
                            "questions.", "pull key operational metrics from your internal systems into one screen"),
    "Trading analytics": ("Trader, account and risk data needs to be clear and close to real time.",
                          "show account performance, drawdown and risk metrics in real time"),
    "Other": ("Data spread across several tools makes it hard to see the full picture quickly.",
              "bring your key numbers into one live view"),
}


def domain(u):
    u = (u or "").lower().replace("https://", "").replace("http://", "").split("/")[0]
    return u[4:] if u.startswith("www.") else u


def root(d):
    p = d.split(".")
    if len(p) >= 3 and p[-2] in ("co", "com", "org", "net") and len(p[-1]) == 2:
        return ".".join(p[-3:])
    return ".".join(p[-2:])


def clean_angle(a):
    a = a.strip()
    for cut in (", so ", " and may need", " — ", "; "):
        if cut in a:
            a = a.split(cut)[0]
    a = a.rstrip(" .,")
    return a + "."


def first_name(n):
    n = re.sub(r"^(dr|mr|mrs|ms)\.?\s+", "", (n or "").strip(), flags=re.I)
    return n.split()[0] if n else ""


def build_email(r, name):
    company = r["Company Name"]
    problem = r["Likely Business Problem"] if r["Likely Business Problem"] in PAIN else "Other"
    pain, benefit = PAIN[problem]
    solution = r["Potential Dashboard Solution"].strip() or "custom reporting dashboard"
    greet = f"Hi {first_name(name)}," if name else f"Hi {company} team,"
    subject = f"{solution} for {company}"
    body = (
        f"{greet}\n\n"
        f"{clean_angle(r['Personalization Angle'])}\n\n"
        f"{pain} I build custom dashboards and reporting tools with React, Next.js and Python that connect "
        f"directly to platforms like HubSpot, Salesforce, Shopify, Stripe, n8n, Make and custom APIs, so the "
        f"numbers update automatically in one place.\n\n"
        f"For {company}, a {solution} could {benefit}. I am happy to share a quick mockup based on your "
        f"current stack, with no obligation.\n\n"
        f"Would you be open to a short call next week to see if this would be useful? If the timing is not "
        f"right, no problem at all.\n\n"
        f"Best regards,\n{SENDER[0]}\n{SENDER[1]}\n{SENDER[2]}"
    )
    return subject, body


def main():
    rows = list(csv.DictReader(open(os.path.join(ROOT, "Dashboard_Development_Prospects_700.csv"), encoding="utf-8-sig")))
    enrich = {}
    for f in sorted(glob.glob(os.path.join(HERE, "enrichment", "*.json"))):
        for e in json.load(open(f, encoding="utf-8")):
            if e.get("email"):
                enrich.setdefault(e["lead_id"], e)

    stats = {"invalid_removed": [], "duplicates_removed": [], "from_enrichment": 0, "from_original": 0}
    seen = set()
    ready, pending = [], []
    for r in rows:
        name = r["Contact Name"].strip()
        email, source, how = r["Business Email"].strip().lower(), r["Email Source"].strip(), r["Email Status"]
        e = enrich.get(r["Lead ID"])
        if not email and e:
            email, source = e["email"].strip().lower().rstrip("."), e.get("source_url", "").strip()
            on_site = root(domain(source)) in (root(domain(r["Website"])), root(email.split("@")[-1]))
            how = ("Found in public search index during enrichment (company page)" if on_site else
                   "Found in public search index during enrichment (third-party source, verify before sending)")
            if not name and e.get("contact_name"):
                name = e["contact_name"].strip()
                r["Job Title"] = e.get("job_title", "")
        notes = [f"Lead {r['Lead ID']}", r["Lead Category"], f"Quality: {r['Lead Quality']}",
                 f"Evidence: {r['Evidence URL']}"]
        if r["Job Title"] and name:
            notes.append(f"Contact title: {r['Job Title']}")
        if email:
            ed = email.split("@")[-1]
            problem = None
            if not EMAIL_RE.match(email) or PLACEHOLDER.search(email):
                problem = "invalid or placeholder"
            elif ed in CONSUMER:
                problem = "consumer domain"
            elif not source.startswith("http"):
                problem = "no source URL"
            elif r["Website"] and root(ed) != root(domain(r["Website"])) and "alias" not in (e or {}).get("note", "").lower():
                problem = "domain does not match company website"
            if problem:
                stats["invalid_removed"].append((r["Company Name"], email, problem))
                email = ""
            elif email in seen:
                stats["duplicates_removed"].append((r["Company Name"], email))
                email = ""
        if email:
            seen.add(email)
            stats["from_enrichment" if e and how.startswith("Found") else "from_original"] += 1
            notes.append(f"Email status: {how}")
            if e and how.startswith("Found") and e.get("note"):
                notes.append(f"Email note: {e['note']}")
        else:
            source = ""
        subject, body = build_email(r, name)
        assert not BAD_TEXT.search(subject + body), (r["Company Name"], subject)
        out = [name, r["Company Name"], r["Website"], email, subject, body, "",
               r["Country"], r["Industry"], source or r["Evidence URL"], " | ".join(n for n in notes if n)]
        (ready if email else pending).append(out)

    wb = Workbook()
    for title, data in (("Send Ready", ready), ("Needs Email", pending)):
        ws = wb.active if title == "Send Ready" else wb.create_sheet(title)
        ws.title = title
        ws.append(COLS)
        for row in data:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
        widths = [18, 28, 30, 32, 45, 80, 10, 14, 22, 45, 60]
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            row[5].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    out = os.path.join(ROOT, "Dashboard_Outreach_Sender.xlsx")
    wb.save(out)
    with open(os.path.join(ROOT, "Dashboard_Outreach_Sender_SendReady.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLS)
        w.writerows(ready)

    report = {
        "Total prospects": len(ready) + len(pending),
        "Prospects with verified/public emails": len(ready),
        "  - emails from original research": stats["from_original"],
        "  - emails from enrichment pass": stats["from_enrichment"],
        "Prospects without emails": len(pending),
        "Duplicate emails removed": len(stats["duplicates_removed"]),
        "Invalid emails removed": len(stats["invalid_removed"]),
    }
    json.dump({"report": report, "invalid": stats["invalid_removed"], "duplicates": stats["duplicates_removed"]},
              open(os.path.join(HERE, "outreach_qc.json"), "w"), indent=2)
    print(json.dumps(report, indent=2))
    for x in stats["invalid_removed"]:
        print("invalid:", x)
    for x in stats["duplicates_removed"]:
        print("duplicate:", x)


if __name__ == "__main__":
    main()
